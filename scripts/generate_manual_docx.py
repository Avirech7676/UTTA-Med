"""Script to compile the complete Research and Development Manual into a publication-grade Word document (.docx)."""

import os
import shutil
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_header(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    if level == 1:
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        run.font.size = Pt(15)
        run.font.color.rgb = RGBColor(27, 54, 93) # Deep Navy
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="8" w:space="4" w:color="1B365D"/></w:pBdr>')
        p._p.get_or_add_pPr().append(pBdr)
    elif level == 2:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run.font.size = Pt(12.5)
        run.font.color.rgb = RGBColor(46, 97, 113) # Teal
    elif level == 3:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(60, 60, 60)
    return p

def add_paragraph_formatted(doc, text, bold_prefix=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.color.rgb = RGBColor(30, 30, 30)
    if text:
        # Simple inline formatting parsing (**bold** and *italic*)
        parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                r = p.add_run(part[2:-2])
                r.bold = True
                r.font.color.rgb = RGBColor(30, 30, 30)
            elif part.startswith('*') and part.endswith('*'):
                r = p.add_run(part[1:-1])
                r.italic = True
                r.font.color.rgb = RGBColor(40, 40, 40)
            else:
                r = p.add_run(part)
                r.font.color.rgb = RGBColor(45, 45, 45)
    return p

def add_callout_box(doc, text, title="NOTE", border_color="1B365D", fill_color="F0F4F8"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, fill_color)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="20" w:space="0" w:color="{border_color}"/>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r_title = p.add_run(f"[{title}] ")
    r_title.bold = True
    r_title.font.size = Pt(9.5)
    r_title.font.color.rgb = RGBColor(27, 54, 93)
    
    r_txt = p.add_run(text)
    r_txt.font.size = Pt(9.5)
    r_txt.font.color.rgb = RGBColor(40, 40, 40)
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_str):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=90, bottom=90, left=140, right=140)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="12" w:space="0" w:color="94A3B8"/>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(code_str)
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(30, 41, 59)
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(4)

def add_table_data(doc, headers, rows_data):
    num_cols = len(headers)
    tbl = doc.add_table(rows=1, cols=num_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color="CBD5E1", sz="4")
    
    # Header
    hdr_cells = tbl.rows[0].cells
    for i, title in enumerate(headers):
        cell = hdr_cells[i]
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, top=100, bottom=100, left=80, right=80)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(title.strip())
        run.bold = True
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(255, 255, 255)
        
    trPr = tbl.rows[0]._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    
    for r_idx, row in enumerate(rows_data):
        row_cells = tbl.add_row().cells
        trPr = tbl.rows[-1]._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        bg_color = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
        for c_idx, val in enumerate(row):
            if c_idx < num_cols:
                cell = row_cells[c_idx]
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.1
                run = p.add_run(val.strip())
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(40, 40, 40)
                
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(4)

def parse_markdown_to_docx(md_path, docx_path):
    print(f"Reading markdown from {md_path}...")
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()
        
    doc = docx.Document()
    
    # Configure 1-inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    # Set default style font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(10.5)
    font.color.rgb = RGBColor(40, 40, 40)
    
    # Title Cover Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(24)
    title_p.paragraph_format.space_after = Pt(4)
    r_title = title_p.add_run("UTTA-Med: Research & Development Manual")
    r_title.bold = True
    r_title.font.size = Pt(22)
    r_title.font.color.rgb = RGBColor(27, 54, 93)
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(14)
    r_sub = sub_p.add_run("Selective Test-Time Adaptation, Mathematical Gating Equivalence, and Runtime Safeguarding under Cross-Hospital Histopathology Domain Shift")
    r_sub.font.size = Pt(12)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(70, 70, 70)
    
    meta_tbl = doc.add_table(rows=1, cols=1)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = meta_tbl.cell(0, 0)
    c.width = Inches(6.5)
    set_cell_background(c, "F1F5F9")
    set_cell_margins(c, top=140, bottom=140, left=180, right=180)
    
    mp = c.paragraphs[0]
    mp.paragraph_format.space_before = Pt(0)
    mp.paragraph_format.space_after = Pt(0)
    mp.paragraph_format.line_spacing = 1.2
    
    runs_data = [
        ("Lead Researcher: ", True, RGBColor(27, 54, 93)),
        ("Avinash Reddy Cheerapareddy\n", False, RGBColor(40, 40, 40)),
        ("Institutional Affiliation: ", True, RGBColor(27, 54, 93)),
        ("Department of Computer Science and Engineering, Medical Imaging Laboratory\n", False, RGBColor(40, 40, 40)),
        ("Software Framework: ", True, RGBColor(27, 54, 93)),
        ("SafeAdapt-TTA / UTTA-Med v2 (Camelyon17-WILDS ResNet-18 Benchmark)\n", False, RGBColor(40, 40, 40)),
        ("Canonical Checkpoint (SHA-256): ", True, RGBColor(27, 54, 93)),
        ("f0b85ec1a89707a8ad1807a544280ac2302a572abd0ba84233dff6d4dd35d000\n", False, RGBColor(80, 80, 80)),
        ("Automated Test Suite Status: ", True, RGBColor(27, 54, 93)),
        ("120 Passed, 2 Skipped, 0 Failed (100% Passing Rate in 11.16s)\n", False, RGBColor(22, 101, 52)),
        ("Numerical Table Audit: ", True, RGBColor(27, 54, 93)),
        ("54 of 54 Bit-Exact Checks Passed (100% Passing Rate)\n", False, RGBColor(22, 101, 52)),
        ("Publication Date: ", True, RGBColor(27, 54, 93)),
        ("October 2026\n", False, RGBColor(40, 40, 40))
    ]
    for text_val, is_b, col in runs_data:
        r = mp.add_run(text_val)
        r.bold = is_b
        r.font.size = Pt(9.5)
        r.font.color.rgb = col
        
    doc.add_page_break()
    
    # State machine for markdown parsing
    in_code = False
    code_lines = []
    in_table = False
    table_headers = []
    table_rows = []
    
    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\r\n')
        
        # Check code block
        if line.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                in_code = False
                code_str = "\n".join(code_lines)
                add_code_block(doc, code_str)
            i += 1
            continue
            
        if in_code:
            code_lines.append(line)
            i += 1
            continue
            
        # Check markdown table
        if "|" in line and line.strip().startswith("|") and line.strip().endswith("|"):
            cells = [c.strip() for c in line.strip().split("|")[1:-1]]
            if not in_table:
                in_table = True
                table_headers = cells
                table_rows = []
                # Next line might be separator |:---|:---|
                if i + 1 < len(lines) and "---" in lines[i+1]:
                    i += 1 # skip separator
            else:
                table_rows.append(cells)
            i += 1
            continue
        else:
            if in_table:
                in_table = False
                if table_headers and table_rows:
                    add_table_data(doc, table_headers, table_rows)
                table_headers = []
                table_rows = []
                
        # Skip pure divider lines
        if line.strip() in ["---", "===", "___"]:
            i += 1
            continue
            
        # Empty line
        if not line.strip():
            i += 1
            continue
            
        # Headings
        if line.startswith("# "):
            add_header(doc, line[2:].strip(), 1)
        elif line.startswith("## "):
            add_header(doc, line[3:].strip(), 2)
        elif line.startswith("### "):
            add_header(doc, line[4:].strip(), 3)
        elif line.startswith("#### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(line[5:].strip())
            r.bold = True
            r.font.size = Pt(10.5)
            r.font.color.rgb = RGBColor(27, 54, 93)
        elif line.startswith("> [!"):
            # Alert / Callout
            alert_type = "NOTE"
            if "IMPORTANT" in line:
                alert_type = "IMPORTANT"
            elif "WARNING" in line or "CAUTION" in line:
                alert_type = "WARNING"
            elif "TIP" in line:
                alert_type = "TIP"
            callout_text = line.split("]", 1)[-1].strip()
            # see if subsequent lines are blockquote
            j = i + 1
            while j < len(lines) and lines[j].startswith(">"):
                callout_text += " " + lines[j][1:].strip()
                j += 1
            add_callout_box(doc, callout_text, title=alert_type)
            i = j - 1
        elif line.startswith(">"):
            add_callout_box(doc, line[1:].strip(), title="RESEARCH NOTE")
        elif line.startswith("- [x]") or line.startswith("- [ ]"):
            status_box = "[PASS] " if line.startswith("- [x]") else "[TODO] "
            add_paragraph_formatted(doc, line[5:].strip(), bold_prefix=status_box, space_after=2)
        elif line.strip().startswith("* ") or line.strip().startswith("- "):
            add_paragraph_formatted(doc, line.strip()[2:], bold_prefix="• ", space_after=3)
        elif re.match(r'^\d+\.\s', line.strip()):
            m = re.match(r'^(\d+\.)\s(.*)', line.strip())
            add_paragraph_formatted(doc, m.group(2), bold_prefix=f"{m.group(1)} ", space_after=3)
        else:
            add_paragraph_formatted(doc, line, space_after=4)
            
        i += 1
        
    if in_table and table_headers and table_rows:
        add_table_data(doc, table_headers, table_rows)
        
    print(f"Saving compiled Word document to {docx_path}...")
    doc.save(docx_path)
    print("Done!")

if __name__ == "__main__":
    src_md = os.path.abspath("docs/final/RESEARCH_AND_DEVELOPMENT_MANUAL.md")
    dest_docx = os.path.abspath("RESEARCH_AND_DEVELOPMENT_MANUAL.docx")
    downloads_docx = os.path.expanduser(r"~\Downloads\RESEARCH_AND_DEVELOPMENT_MANUAL.docx")
    
    parse_markdown_to_docx(src_md, dest_docx)
    
    try:
        shutil.copyfile(dest_docx, downloads_docx)
        print(f"Successfully copied Word document to Downloads: {downloads_docx}")
    except Exception as e:
        print(f"Could not copy to Downloads: {e}")
